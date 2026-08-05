import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_callout(doc, text, title=""):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "1E2636")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    if title:
        run_t = p.add_run(f"📌 {title}\n")
        run_t.bold = True
        run_t.font.color.rgb = RGBColor(88, 166, 255)
        run_t.font.size = Pt(10.5)
    run_txt = p.add_run(text)
    run_txt.font.size = Pt(9.5)
    run_txt.font.color.rgb = RGBColor(201, 209, 217)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def build_document():
    doc = docx.Document()
    
    # Set standard page margins (1 inch all around)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    # Styles Setup
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Arial'
    normal_style.font.size = Pt(10.5)
    normal_style.font.color.rgb = RGBColor(36, 41, 47)

    # -------------------------------------------------------------
    # TITLE SECTION
    # -------------------------------------------------------------
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("JANUS SOC & NOC PLATFORM")
    r_title.bold = True
    r_title.font.size = Pt(24)
    r_title.font.color.rgb = RGBColor(15, 23, 42)
    p_title.paragraph_format.space_after = Pt(2)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("Comprehensive System Architecture & Module Reference Manual")
    r_sub.font.size = Pt(13)
    r_sub.font.color.rgb = RGBColor(100, 116, 139)
    p_sub.paragraph_format.space_after = Pt(24)

    add_callout(
        doc,
        "This document contains detailed functional, technical, and operational documentation for all integrated modules of the JANUS Security Operations Center (SOC) & Network Operations Center (NOC) Telemetry Platform. It outlines core monitoring capabilities, SIEM log parsing pipelines, network tools, and security auditing features.",
        "EXECUTIVE OVERVIEW"
    )

    # -------------------------------------------------------------
    # SECTION 1: ARCHITECTURE OVERVIEW
    # -------------------------------------------------------------
    h1 = doc.add_heading("1. System Architecture & Tech Stack", level=1)
    h1.runs[0].font.color.rgb = RGBColor(30, 41, 59)
    
    p = doc.add_paragraph(
        "JANUS is built on a high-performance, containerized microservices architecture designed to collect, process, index, and visualize network and security events in real time. The platform consists of three primary containerized services orchestrated via Docker Compose:"
    )
    
    table_stack = doc.add_table(rows=1, cols=4)
    table_stack.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table_stack.rows[0].cells
    headers = ["Service Container", "Technology Stack", "Port / Protocol", "Primary Purpose"]
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        hdr_cells[i].paragraphs[0].runs[0].bold = True
        hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        set_cell_background(hdr_cells[i], "0F172A")
        set_cell_margins(hdr_cells[i], 100, 100, 120, 120)

    stack_data = [
        ("janus-web-frontend", "PHP 8.1 / Apache / JavaScript", "Port 8000 (HTTP)", "Web User Interface, RESTful AJAX endpoints, telemetry dashboards, and report engines."),
        ("janus-syslog-receiver", "Python 3.10 / PyMySQL (Host Net)", "Port 5140 (UDP/TCP)", "Multi-threaded SIEM daemon that receives RFC 3164 syslogs, extracts PRI facility/severity, and writes to MySQL."),
        ("janus-mysql-db", "MySQL 8.0 Enterprise Relational DB", "Port 3307 (Internal 3306)", "Central database storing 35 tables & 5 analytical views covering accounts, printers, syslog entries, and rollups.")
    ]

    for c1, c2, c3, c4 in stack_data:
        row_cells = table_stack.add_row().cells
        for i, val in enumerate([c1, c2, c3, c4]):
            row_cells[i].text = val
            set_cell_margins(row_cells[i], 80, 80, 100, 100)
            if i == 0:
                row_cells[i].paragraphs[0].runs[0].bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------
    # SECTION 2: MODULE DETAILS
    # -------------------------------------------------------------
    h2 = doc.add_heading("2. Core System Modules & Features", level=1)
    h2.runs[0].font.color.rgb = RGBColor(30, 41, 59)

    modules_info = [
        {
            "num": "2.1",
            "name": "Main NOC Dashboard",
            "file": "home.php",
            "category": "Core NOC Monitoring",
            "desc": "The primary operational control center providing high-level situational awareness across network nodes, uptime health, active incidents, and bandwidth utilization.",
            "features": [
                "Real-time visual indicator widgets for total configured devices, online status, critical alerts, and active tickets.",
                "Global network availability meter and link health summary.",
                "Quick-action links to inventory, topology maps, and threat detection modules."
            ]
        },
        {
            "num": "2.2",
            "name": "Device Inventory & Asset Catalog",
            "file": "manage.php",
            "category": "Core NOC Monitoring",
            "desc": "Centralized database management interface for all monitored network infrastructure, including routers, switches, firewalls, servers, and printers.",
            "features": [
                "Full CRUD capabilities for network assets with custom fields for IP, vendor, model, city, sub-office, and location tags.",
                "SNMP community string (v1, v2c, v3) and credential management per device.",
                "Automatic active status detection and single-click manual SNMP poll triggers."
            ]
        },
        {
            "num": "2.3",
            "name": "Real-Time Ping & Health Telemetry",
            "file": "monitor.php",
            "category": "Core NOC Monitoring",
            "desc": "Active polling service that continuously monitors ICMP reachability, round-trip latency (ms), and packet loss percentages for all registered IP addresses.",
            "features": [
                "Background ICMP ping execution with customizable polling intervals (default 120s).",
                "Historical latency tracking stored in ping_logs and ping_logs_archive tables.",
                "Visual status indicators (Online, Degraded, Offline) with automated fault alerts."
            ]
        },
        {
            "num": "2.4",
            "name": "Network Topology Mapping",
            "file": "maps.php",
            "category": "Core NOC Monitoring",
            "desc": "Interactive graphical representation of network interconnectivity, link health, and parent-child infrastructure relationships.",
            "features": [
                "Dynamic HTML5 Canvas / SVG rendering of network nodes and connecting links.",
                "Real-time color-coded link status overlays (Green = UP, Red = DOWN, Yellow = DEGRADED).",
                "Interactive node selection providing instant IP details and device performance summaries."
            ]
        },
        {
            "num": "2.5",
            "name": "Performance Metrics & Resource Monitoring",
            "file": "resources.php",
            "category": "Core NOC Monitoring",
            "desc": "Detailed hardware utilization analytics tracking system resources across network appliances via SNMP queries.",
            "features": [
                "CPU load average, RAM memory consumption, and disk partition usage gauges.",
                "Interface bandwidth utilization graphs (Inbound / Outbound Bps).",
                "Historical resource trending graphs with customizable reporting windows."
            ]
        },
        {
            "num": "2.6",
            "name": "IP Address Management (IPAM)",
            "file": "modules/ipam/ipam.php",
            "category": "NOC Tools",
            "desc": "Comprehensive IP space management module tracking IPv4/IPv6 subnet allocations, address assignments, and conflict detection.",
            "features": [
                "Subnet visualizer displaying assigned, reserved, and available IP addresses in color-coded blocks.",
                "IP status tracking (Static, DHCP, Reserved, Available) and gateway binding.",
                "Automated detection of unassigned IP addresses appearing in syslog streams."
            ]
        },
        {
            "num": "2.7",
            "name": "Layer 2 & Port Security / NAC",
            "file": "modules/nac/nac.php",
            "category": "NOC Tools",
            "desc": "Network Access Control module monitoring managed switch ports, MAC address tables, and enforcement policies.",
            "features": [
                "Switch port mapping displaying MAC address to switch-port associations.",
                "Unauthorized device detection and rogue MAC isolation capability.",
                "Remote port administrative status toggling (Enable / Disable port via SNMP)."
            ]
        },
        {
            "num": "2.8",
            "name": "Printer Fleet Management & Telemetry",
            "file": "modules/printers/printers.php & ss/printer_activity.php",
            "category": "NOC Tools / Telemetry",
            "desc": "Advanced multi-vendor printer audit and telemetry platform supporting HP, Ricoh, Canon, Xerox, Lexmark, Kyocera, Brother, and Epson printers.",
            "features": [
                "Automated classification of syslog events into Print Jobs, Scan Jobs (including Save to USB, Scan to Email/SMB), Copy Jobs, Authentication, Admin & Config, and System & Hardware events.",
                "HP-specific key-value parser handling domain user stripping (e.g. RELIANCE\\user -> user), source_IP extraction, and job outcome status.",
                "Distinct color-coded status badges: Green (PRINTED/SUCCESS), Cyan (SCANNED), Purple (COPIED), Red (FAILED), Orange (WARNING/MODIFIED), Slate (COMPLETED).",
                "Low-power state, paper jam, toner low, and error cleared telemetry tracking with real raw syslog viewer drawer."
            ]
        },
        {
            "num": "2.9",
            "name": "Security Analytics Console",
            "file": "reports.php",
            "category": "SOC Security Operations",
            "desc": "Centralized SIEM security dashboard visualizing network threat activity, authentication anomalies, and security policy violations.",
            "features": [
                "Aggregated security event timeline with severity distribution metrics.",
                "Tabbed drilldown interface for Threat Hunting, Brute-Force, DDoS, and Anomaly Detection.",
                "Export options (CSV / PDF / Report Generation) and scheduled report delivery."
            ]
        },
        {
            "num": "2.10",
            "name": "Traffic Analyzer & NetFlow Analyzer",
            "file": "fflow.php & ss/traffic.php",
            "category": "SOC Security Operations",
            "desc": "Deep network traffic analysis module aggregating flow logs, application usage, and bandwidth metrics.",
            "features": [
                "Protocol breakdown (HTTP, HTTPS, SSH, DNS, SMB, RDP, MySQL, etc.).",
                "Top talker identification by Source IP, Destination IP, and Volume (Bytes/Packets).",
                "Hourly and daily rollup aggregations powered by syslog_traffic_daily and syslog_traffic_hourly tables."
            ]
        },
        {
            "num": "2.11",
            "name": "IPS Logs & Intrusion Detection",
            "file": "ips_logs.php",
            "category": "SOC Security Operations",
            "desc": "Dedicated Intrusion Prevention System (IPS) log audit interface parsing firewall and sensor security events.",
            "features": [
                "Signature ID (SID) and attack name classification.",
                "Threat severity badges (Critical, High, Medium, Low) and action tracking (Dropped, Blocked, Passed).",
                "One-click IP drilldown to inspect all security telemetry associated with an attacking host."
            ]
        },
        {
            "num": "2.12",
            "name": "Vulnerability Scanner",
            "file": "modules/scanners/vuln_scan.php",
            "category": "SOC Security Operations",
            "desc": "Automated network service and vulnerability scanning module evaluating node exposure across open ports.",
            "features": [
                "Nmap / Socket-based port scanning against targeted subnets or single IP hosts.",
                "Service version detection (Apache, SSH, OpenSSL, MySQL, RDP).",
                "CVE vulnerability mapping and remediation guidance generation."
            ]
        },
        {
            "num": "2.13",
            "name": "Appliance Configuration Backup",
            "file": "modules/config_backup/",
            "category": "Appliance Management",
            "desc": "Automated configuration backup and version control module for network switches, routers, and firewalls.",
            "features": [
                "Scheduled SSH/Telnet/SNMP configuration pulling (running-config / startup-config).",
                "Side-by-side diff comparison highlighting configuration drift between backup revisions.",
                "One-click rollback script generation."
            ]
        },
        {
            "num": "2.14",
            "name": "Policy & Compliance Auditing",
            "file": "compliance.php",
            "category": "Appliance Management",
            "desc": "Security compliance evaluation engine assessing infrastructure configurations against established hardening standards.",
            "features": [
                "Compliance score calculation based on enabled security features (SNMPv3, SSH-only, Syslog enabled, Strong Passwords).",
                "Detailed pass/fail breakdown for ISO 27001, CIS Benchmarks, and internal security baselines.",
                "Remediation checklist for non-compliant appliances."
            ]
        },
        {
            "num": "2.15",
            "name": "SIEM Log Management & Syslog Receiver",
            "file": "modules/logmanage/ & syslog_receiver.py",
            "category": "SIEM Engine",
            "desc": "Core log ingest engine capable of handling high-volume syslog traffic from firewalls, routers, switches, and printers.",
            "features": [
                "Multi-threaded UDP/TCP listening daemon running on host network mode (port 5140 / 514).",
                "RFC 3164 PRI parsing: extracts Facility (0-23) and Severity (0-7: Emergency to Debug) from <PRI> headers.",
                "Automated database connection auto-reconnect (db.ping(reconnect=True)) preventing broken pipes during idle periods.",
                "Raw log file logging to /var/log/janus_siem/ per device IP and date."
            ]
        },
        {
            "num": "2.16",
            "name": "User Management & Role-Based Access Control (RBAC)",
            "file": "users.php & auth_check.php",
            "category": "Administration",
            "desc": "Identity management module securing access to JANUS features based on assigned user roles.",
            "features": [
                "Role definitions: Administrator (Full Access), Analyst (Operational & Security Access), Backup Manager (Configs).",
                "Bcrypt password hashing ($2y$10$) and secure session management.",
                "Fine-grained page level permission checks via nac_can_access() function."
            ]
        }
    ]

    for m in modules_info:
        h_mod = doc.add_heading(f"{m['num']} {m['name']}", level=2)
        h_mod.runs[0].font.color.rgb = RGBColor(30, 41, 59)

        p_meta = doc.add_paragraph()
        r_f = p_meta.add_run("Primary Script: ")
        r_f.bold = True
        p_meta.add_run(f"{m['file']}  |  ")
        r_c = p_meta.add_run("Category: ")
        r_c.bold = True
        p_meta.add_run(m['category'])
        p_meta.paragraph_format.space_after = Pt(4)

        p_desc = doc.add_paragraph(m['desc'])
        p_desc.paragraph_format.space_after = Pt(6)

        p_feat_h = doc.add_paragraph()
        r_fh = p_feat_h.add_run("Key Features & Functionality:")
        r_fh.bold = True
        p_feat_h.paragraph_format.space_after = Pt(2)

        for feat in m['features']:
            bp = doc.add_paragraph(style='List Bullet')
            bp.paragraph_format.space_before = Pt(1)
            bp.paragraph_format.space_after = Pt(2)
            bp.add_run(feat)

        doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Save Document
    target_path = "/home/tamseel-hassan/Desktop/JANUS_System_Modules_Documentation.docx"
    doc.save(target_path)
    print(f"Document successfully created at: {target_path}")

if __name__ == '__main__':
    build_document()
